"""
Simulate an email with a .docx attachment and route it through the VCF
intake router. This tests the email-intake attachment path without a live
email provider.

Run from repo root:
    python scripts/test_email_docx_route.py
"""
import sys
from email.message import EmailMessage
from email.policy import default
from io import BytesIO
from pathlib import Path

from docx import Document

# Ensure repo root is on path
ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))

from backend.demo1.intake_router import route_file


def build_docx() -> bytes:
    doc = Document()
    doc.add_heading("VCF Medical Record", level=1)
    doc.add_paragraph("Patient: Chen Weiming")
    doc.add_paragraph("DOB: 1958-03-12")
    doc.add_paragraph("Diagnosis: Stage IIIA lung adenocarcinoma (EGFR exon 19 deletion)")
    doc.add_paragraph("WTC Health Program registered: 2018")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def main() -> None:
    docx_bytes = build_docx()

    msg = EmailMessage(policy=default)
    msg["From"] = "dr.office@example.com"
    msg["To"] = "vcfclaim@wawvcf.com"
    msg["Subject"] = "Medical records for Chen Weiming"
    msg.set_content("Please find attached medical records for the VCF claim.")
    msg.add_attachment(
        docx_bytes,
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename="chen_weiming_medical_records.docx",
    )

    attachments = []
    for part in msg.walk():
        if part.get_content_disposition() == "attachment":
            filename = part.get_filename()
            data = part.get_payload(decode=True)
            if data:
                attachments.append({"filename": filename, "data": data})

    print(f"Attachments found: {len(attachments)}")
    for att in attachments:
        print(f"  {att['filename']}: {len(att['data'])} bytes")

    for att in attachments:
        result = route_file(att["filename"], att["data"], "waw_vcf")
        print("\n--- Routing result ---")
        print(f"filename    : {att['filename']}")
        print(f"route_taken : {result.get('route_taken')}")
        print(f"engine      : {result.get('engine')}")
        print(f"word_count  : {result.get('word_count')}")
        print(f"confidence  : {result.get('confidence')}")
        print(f"text preview: {result.get('text', '')[:400]}")


if __name__ == "__main__":
    main()
